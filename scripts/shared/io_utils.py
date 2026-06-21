"""Reusable I/O helpers for COSMET-format measurement files."""
from __future__ import annotations

import csv
import re
from pathlib import Path

import pandas as pd

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None  # type: ignore[assignment,misc]


# ---------------------------------------------------------------------------
# Extract CFP
# ---------------------------------------------------------------------------
def extract_cfp_from_xlsx(path: Path) -> int | None:
    """Read total CFP from a COSMET-format measurement XLSX."""
    try:
        df = pd.read_excel(path, engine="openpyxl")
        cfp_col = df.columns[df.columns.str.upper().str.contains("CFP")].tolist()
        if not cfp_col:
            return None
        total_rows = (
            df[df.iloc[:, -2].astype(str).str.upper() == "TOTAL"]
            if df.shape[1] >= 2
            else pd.DataFrame()
        )
        if not total_rows.empty:
            return int(total_rows.iloc[0][cfp_col[0]])
        valid = pd.to_numeric(df[cfp_col[0]], errors="coerce").dropna()
        return int(valid.sum()) if not valid.empty else None
    except Exception:
        return None


def extract_cfp_from_csv(path: Path) -> int | None:
    """Read total CFP from a COSMET-format measurement CSV."""
    try:
        df = pd.read_csv(path)
        dm_col = [c for c in df.columns if "DM" in c.upper()]
        cfp_col = [c for c in df.columns if "CFP" in c.upper()]
        if not cfp_col:
            return None
        total_row = (
            df[df[dm_col[0]].astype(str).str.upper() == "TOTAL"]
            if dm_col
            else pd.DataFrame()
        )
        if not total_row.empty:
            return int(total_row.iloc[0][cfp_col[0]])
        valid = pd.to_numeric(df[cfp_col[0]], errors="coerce").dropna()
        return int(valid.sum()) if not valid.empty else None
    except Exception:
        return None


def extract_movements_from_csv(path: Path) -> dict | None:
    """Extract per-sentence E/R/W/X counts from a COSMET CSV."""
    try:
        df = pd.read_csv(path)
        dm_col = [c for c in df.columns if "DM" in c.upper()][0]
        movements = df[dm_col].dropna().tolist()
        entry = sum(1 for m in movements if "entry" in str(m).lower())
        read = sum(1 for m in movements if "read" in str(m).lower())
        write = sum(1 for m in movements if "write" in str(m).lower())
        exit_ = sum(
            1
            for m in movements
            if "exit" in str(m).lower() and "must be counted" not in str(m).lower()
        )
        exit_once = sum(
            1 for m in movements if "must be counted only once" in str(m).lower()
        )
        exit_total = min(1, exit_once) + exit_ if exit_once else exit_
        return {"entry": entry, "read": read, "write": write, "exit": exit_total}
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Save measurements
# ---------------------------------------------------------------------------
def save_measurement_xlsx(result: dict, path: Path) -> None:
    """Save COSMET measurement to XLSX (same schema as original COSMET)."""
    if Workbook is None:
        raise ImportError("openpyxl is required for XLSX output")
    wb = Workbook()
    ws = wb.active
    ws.title = "Measurement"
    ws.append(["FU", "Sub-process", "DG", "OOI", "DM", "CFP"])
    t = result["table"]
    for i in range(len(t["FU"])):
        ws.append([t["FU"][i], t["Sub-process"][i], t["DG"][i], t["OOI"][i], t["DM"][i], t["CFP"][i]])

    ws2 = wb.create_sheet(title="Triggering event")
    ws2.append(["Triggering event", result["triggering_event"]])

    ws3 = wb.create_sheet(title="FU")
    ws3.append(["Functional User", result["functional_user"]])

    ws4 = wb.create_sheet(title="Time")
    ws4.append(["DURATION", result["duration"]])

    wb.save(path)


def save_measurement_csv(result: dict, path: Path) -> None:
    """Save COSMET measurement to CSV."""
    t = result["table"]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["FU", "Sub-process", "DG", "OOI", "DM", "CFP"])
        for i in range(len(t["FU"])):
            w.writerow([t["FU"][i], t["Sub-process"][i], t["DG"][i], t["OOI"][i], t["DM"][i], t["CFP"][i]])


# ---------------------------------------------------------------------------
# File discovery
# ---------------------------------------------------------------------------
def find_fp_files(input_dir: Path) -> list[Path]:
    """Find FP prompt .txt files and sort numerically."""
    files = list(input_dir.glob("*.txt"))

    def sort_key(p: Path) -> int:
        m = re.search(r"(\d+)", p.stem)
        return int(m.group(1)) if m else 0

    return sorted(files, key=sort_key)
