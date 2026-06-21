#!/usr/bin/env python3
"""Consolidate all dataset sources into a unified manifest.csv.

Reads every known data source — with and without COSMIC ground truth —
and produces a single CSV at datasets/manifest.csv plus summary stats.

Usage:
    python scripts/consolidate_manifest.py          # default
    python scripts/consolidate_manifest.py --out datasets/manifest.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent  # software-size-estimator/

MANIFEST_SCHEMA = [
    "id",
    "source",
    "format",
    "text",
    "cfp",
    "entry",
    "read",
    "write",
    "exit",
    "domain",
    "project",
    "has_ground_truth",
]


def _extract_cfp_from_xlsx(path: Path) -> int:
    """Extract total CFP from a COSMET ground truth xlsx (last row TOTAL)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    total = 0
    for row in ws.iter_rows(values_only=True):
        if row and row[-2] == "TOTAL":
            try:
                total = int(row[-1])
            except (ValueError, TypeError):
                total = 0
            break
    wb.close()
    return total


def _read_fp_text(path: Path) -> str:
    """Read the full text of an FP_prompts/*.txt file."""
    try:
        return path.read_text(encoding="utf-8").strip()
    except Exception:
        return ""


def load_zenodo(rows: list[dict]) -> None:
    """Zenodo heterogeneous dataset — 2,236 sentences with COSMIC GT."""
    src = ROOT / "datasets" / "raw" / "zenodo-heterogeneous" / "heterogeneous_dataset.csv"
    if not src.exists():
        print(f"  [SKIP] {src} not found")
        return
    df = pd.read_csv(src)
    for _, r in df.iterrows():
        rows.append(
            {
                "id": f"ZENODO-{_ + 1:04d}",
                "source": "zenodo-heterogeneous",
                "format": "sentence",
                "text": str(r.get("requirement_en", "")),
                "cfp": int(r.get("CFP", 0)),
                "entry": int(r.get("entry", 0)),
                "read": int(r.get("read", 0)),
                "write": int(r.get("write", 0)),
                "exit": int(r.get("exit", 0)),
                "domain": "mixed",
                "project": "zenodo",
                "has_ground_truth": True,
            }
        )
    print(f"  Zenodo: {len(df)} sentences loaded")


def load_ochodek(rows: list[dict]) -> None:
    """Ochodek 437 UC dataset — use-case names with CFP GT."""
    src = ROOT / "datasets" / "raw" / "ochodek-uc-437" / "use-cases.csv"
    if not src.exists():
        print(f"  [SKIP] {src} not found")
        return
    df = pd.read_csv(src)
    for _, r in df.iterrows():
        rows.append(
            {
                "id": f"OCHODEK-{r['ProjectID']}-{r['UC']}",
                "source": "ochodek-uc-437",
                "format": "use_case_name",
                "text": str(r.get("TitleTokens", "")),
                "cfp": int(r.get("Cfp", 0)),
                "entry": -1,
                "read": -1,
                "write": -1,
                "exit": -1,
                "domain": "MIS",
                "project": str(r.get("ProjectID", "")),
                "has_ground_truth": True,
            }
        )
    print(f"  Ochodek: {len(df)} use-case names loaded")


def load_cosmet_devito(rows: list[dict]) -> None:
    """COSMET/De Vito datasets — use cases with xlsx ground truth."""
    base = ROOT / "replications" / "cosmet-devito" / "datasets"
    systems = [
        ("MIS", "ALBERGATE", "MIS"),
        ("IoT_Telemedicine", "K01719", "IoT"),
        ("IoT_Telemedicine", "K01720", "IoT"),
        ("IoT_Telemedicine", "K01726", "IoT"),
        ("ML", "U-CURE", "ML"),
        ("Real-time", "RiceCooker", "Real-time"),
        ("Real-time", "AutomaticLineSwitching", "Real-time"),
    ]
    total = 0
    for domain_dir, system, domain_label in systems:
        fp_dir = base / domain_dir / system / "FP_prompts"
        gt_dir = base / domain_dir / system / "ground truth"
        if not fp_dir.exists():
            continue
        fp_files = sorted(fp_dir.glob("FP*.txt"), key=lambda p: int(p.stem[2:]))
        for fp_file in fp_files:
            fp_num = fp_file.stem[2:]
            text = _read_fp_text(fp_file)
            cfp = -1
            has_gt = False
            gt_xlsx = gt_dir / f"FP{fp_num}.xlsx"
            if gt_xlsx.exists():
                cfp = _extract_cfp_from_xlsx(gt_xlsx)
                has_gt = True
            rows.append(
                {
                    "id": f"COSMET-{system}-FP{fp_num}",
                    "source": f"cosmet-devito/{system}",
                    "format": "use_case",
                    "text": text[:500] + ("..." if len(text) > 500 else ""),
                    "cfp": cfp,
                    "entry": -1,
                    "read": -1,
                    "write": -1,
                    "exit": -1,
                    "domain": domain_label,
                    "project": system,
                    "has_ground_truth": has_gt,
                }
            )
            total += 1
    print(f"  COSMET/De Vito: {total} use cases loaded")


def load_fr_nfr(rows: list[dict]) -> None:
    """FR/NFR datasets — requirements without COSMIC GT."""
    manifest = ROOT / "datasets" / "prepared" / "fr-nfr-cosmet" / "xlsx" / "manifest.csv"
    if not manifest.exists():
        print(f"  [SKIP] {manifest} not found")
        return
    df = pd.read_csv(manifest)
    for _, r in df.iterrows():
        rows.append(
            {
                "id": f"FRNFR-{_ + 1:04d}",
                "source": "fr-nfr-xlsx",
                "format": "sentence",
                "text": str(r.get("requirement", "")),
                "cfp": -1,
                "entry": -1,
                "read": -1,
                "write": -1,
                "exit": -1,
                "domain": "mixed",
                "project": "fr-nfr",
                "has_ground_truth": False,
            }
        )
    print(f"  FR/NFR: {len(df)} requirements loaded")


def load_xml_srs(rows: list[dict]) -> None:
    """XML-SRS sections — without COSMIC GT."""
    manifest = ROOT / "datasets" / "prepared" / "xml-cosmet" / "manifest.csv"
    if not manifest.exists():
        print(f"  [SKIP] {manifest} not found")
        return
    df = pd.read_csv(manifest)
    for _, r in df.iterrows():
        rows.append(
            {
                "id": f"XML-{r.get('doc', 'unknown')}-{_ + 1:04d}",
                "source": "xml-requirements",
                "format": "srs",
                "text": str(r.get("section_title", "")),
                "cfp": -1,
                "entry": -1,
                "read": -1,
                "write": -1,
                "exit": -1,
                "domain": "mixed",
                "project": str(r.get("doc", "unknown")),
                "has_ground_truth": False,
            }
        )
    print(f"  XML-SRS: {len(df)} sections loaded")


def print_summary(df: pd.DataFrame) -> None:
    """Print dataset summary statistics."""
    print("\n" + "=" * 60)
    print("DATASET MANIFEST — SUMMARY")
    print("=" * 60)

    print(f"\nTotal items: {len(df)}")
    gt = df[df["has_ground_truth"]]
    no_gt = df[~df["has_ground_truth"]]
    print(f"  With COSMIC ground truth: {len(gt)}")
    print(f"  Without ground truth:     {len(no_gt)}")

    print("\n--- By source ---")
    for src, grp in df.groupby("source"):
        gt_count = grp["has_ground_truth"].sum()
        print(f"  {src:<30s}  N={len(grp):>5d}  GT={gt_count:>5d}")

    print("\n--- By format ---")
    for fmt, grp in df.groupby("format"):
        gt_count = grp["has_ground_truth"].sum()
        print(f"  {fmt:<20s}  N={len(grp):>5d}  GT={gt_count:>5d}")

    print("\n--- By domain ---")
    for dom, grp in df.groupby("domain"):
        gt_count = grp["has_ground_truth"].sum()
        print(f"  {dom:<20s}  N={len(grp):>5d}  GT={gt_count:>5d}")

    print("\n--- CFP distribution (ground truth subset only) ---")
    cfp_gt = gt["cfp"]
    cfp_valid = cfp_gt[cfp_gt >= 0]
    if len(cfp_valid) > 0:
        print(f"  N       = {len(cfp_valid)}")
        print(f"  Mean    = {cfp_valid.mean():.2f}")
        print(f"  Median  = {cfp_valid.median():.1f}")
        print(f"  Std     = {cfp_valid.std():.2f}")
        print(f"  Min     = {cfp_valid.min()}")
        print(f"  Max     = {cfp_valid.max()}")
        print(f"  Q25     = {cfp_valid.quantile(0.25):.1f}")
        print(f"  Q75     = {cfp_valid.quantile(0.75):.1f}")

    print("\n--- Ground truth by source ---")
    for src, grp in gt.groupby("source"):
        cfp_vals = grp["cfp"][grp["cfp"] >= 0]
        if len(cfp_vals):
            print(
                f"  {src:<30s}  N={len(cfp_vals):>5d}  "
                f"CFP mean={cfp_vals.mean():.1f}  "
                f"median={cfp_vals.median():.0f}  "
                f"range=[{cfp_vals.min()}, {cfp_vals.max()}]"
            )

    print()


def main():
    parser = argparse.ArgumentParser(description="Consolidate datasets into unified manifest")
    parser.add_argument(
        "--out",
        default=str(ROOT / "datasets" / "manifest.csv"),
        help="Output CSV path (default: datasets/manifest.csv)",
    )
    args = parser.parse_args()

    print("Loading dataset sources...\n")

    rows: list[dict] = []
    load_zenodo(rows)
    load_ochodek(rows)
    load_cosmet_devito(rows)
    load_fr_nfr(rows)
    load_xml_srs(rows)

    df = pd.DataFrame(rows, columns=MANIFEST_SCHEMA)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_path, index=False, quoting=csv.QUOTE_NONNUMERIC)
    print(f"\nManifest written to {out_path} ({len(df)} rows)")

    print_summary(df)


if __name__ == "__main__":
    main()
